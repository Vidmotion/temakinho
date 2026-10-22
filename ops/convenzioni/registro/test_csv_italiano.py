"""Importazione di un export CSV di Elementor con impostazioni italiane (LibreOffice headless).

Uso: python3 test_csv_italiano.py <registro.xlsx> <cartella lavoro>

Il CSV viene scritto come lo produce Elementor › Submissions › Export (BOM UTF-8, valori codificati
JSON uno per uno: virgolette interne come \\" e barre come \\/), poi importato da LibreOffice con
filtro CSV: separatore virgola, UTF-8, lingua italiano (it-IT), riconoscimento di date e numeri.
I valori importati vengono incollati nel foglio «Importa CSV» di una copia del registro, che viene
ricalcolata. Prova SOLO con LibreOffice: il comportamento di Excel (desktop o Online) va
verificato a parte. Solo dati fittizi; il file consegnato non viene modificato.
"""
import sys, os, json, subprocess, datetime
import openpyxl

SRC, WORK = sys.argv[1], sys.argv[2]
SOFF = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
os.makedirs(WORK, exist_ok=True)
LOG, P, F = [], 0, 0


def ok(desc, cond, got=""):
    global P, F
    P, F = (P + 1, F) if cond else (P, F + 1)
    LOG.append(("OK    | " if cond else "FAIL  | ") + desc + ("" if cond else " → {!r}".format(got)))


def j(v, slashes=True):
    s = json.dumps(v, ensure_ascii=False)
    return s.replace("/", "\\/") if slashes else s


LABELS = ["Nome azienda", "Sito web azienda (facoltativo)", "Email di riferimento", "Numero dipendenti",
          "Settore dell'azienda", "Città di interesse per la convenzione (una o più)", "Privacy Policy", "ID richiesta"]
BASE = ["Form Name (ID)", "Submission ID", "Created At", "User ID", "User Agent", "User IP", "Referrer"]
ROWS = [
    # azienda, sito, email, dipendenti, settore, città, sid, creato
    ("Società Caffè Già – Prova Srl", "https://www.società-prova.example/chi-siamo", "referente@prova.invalid", "1200", "Ristorazione & catering", "Milano, Reggio Calabria", 701, "2026-10-22 09:05:00"),
    ("Ünïcödé Ltd — prova", "", "b@prova.invalid", "7", "Servizi €", "Lione", 702, "2026-01-31 23:59:59"),
    ("Azienda con virgola, Srl", "http://azienda.example", "c@prova.invalid", "35", "Logistica", "Roma, Rimini, Bologna", 703, "2026-03-29 02:30:00"),
    ('Azienda "Alfa" Srl', "", "d@prova.invalid", "12", "Retail", "Milano", 704, "2026-10-22 18:00:00"),
]
lines = [",".join([j(l, slashes=False) for l in LABELS] + BASE)]
for az, sito, em, dip, sett, citta, sid, created in ROWS:
    vals = [j(az), j(sito), j(em), j(dip), j(sett), j(citta), j("on"), j("CONV-20261022-ABCD{:04d}".format(sid))]
    base = [j("Convenzioni aziendali"), j(sid), j(created), j(0), j("Mozilla/5.0 (prova)", slashes=False),
            j("203.0.113.9"), j("https://temakinho.com/convenzioni-aziendali/", slashes=False)]
    lines.append(",".join(vals + base))
csv_path = os.path.join(WORK, "export-elementor.csv")
with open(csv_path, "w", encoding="utf-8", newline="") as fh:
    fh.write("﻿" + "\n".join(lines))

# Importazione con impostazioni italiane: 44=virgola, 34=virgolette, 76=UTF-8, riga 1, lingua 1040 (it-IT),
# campi tra virgolette come testo = no, riconoscimento numeri speciali (date) = sì.
subprocess.run([SOFF, "--headless", "--norestore", "--infilter=CSV:44,34,76,1,,1040,false,true", "--convert-to", "xlsx",
                "--outdir", WORK, csv_path], check=True, capture_output=True)
imp = openpyxl.load_workbook(os.path.join(WORK, "export-elementor.xlsx")).active
grid = [[c.value for c in r] for r in imp.iter_rows()]
ok("CSV: intestazioni lette con accenti", grid[0][5] == LABELS[5], grid[0][5])
ok("CSV: BOM non finito nella prima intestazione", grid[0][0] == "Nome azienda", grid[0][0])
ok("CSV: 4 righe di dati", len(grid) == 5, len(grid))
ok("CSV: accenti nel nome azienda", grid[1][0] == ROWS[0][0], grid[1][0])
ok("CSV: caratteri Unicode e simbolo dell'euro", grid[2][0] == ROWS[1][0] and grid[2][4] == "Servizi €", (grid[2][0], grid[2][4]))
ok("CSV: virgola dentro un valore tra virgolette", grid[3][0] == "Azienda con virgola, Srl", grid[3][0])
ok("CSV: dipendenti riconosciuto come numero (it-IT, nessun separatore delle migliaia)", grid[1][3] == 1200, grid[1][3])
ok("CSV: Submission ID numerico", grid[1][9] == 701, grid[1][9])
ok("CSV: data e ora riconosciute come data", isinstance(grid[1][10], datetime.datetime) and grid[1][10] == datetime.datetime(2026, 10, 22, 9, 5), grid[1][10])
quoted = grid[4][0]
ok("CSV: virgolette interne (Elementor le scrive come \\\") — limite noto dell'export", quoted == 'Azienda "Alfa" Srl', quoted)

# Incolla nel foglio «Importa CSV» di una copia del registro e ricalcola.
wb = openpyxl.load_workbook(SRC)
im = wb["Importa CSV"]
for i, row in enumerate(grid, start=1):
    for k, v in enumerate(row, start=1):
        im.cell(i, k, v)
src2 = os.path.join(WORK, "registro-con-import.xlsx")
wb.save(src2)
out = os.path.join(WORK, "calc")
os.makedirs(out, exist_ok=True)
subprocess.run([SOFF, "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", out, src2], check=True, capture_output=True)
im = openpyxl.load_workbook(os.path.join(out, "registro-con-import.xlsx"), data_only=True)["Importa CSV"]
ok("registro: intestazioni OK", im["AO12"].value == "OK", im["AO12"].value)
esiti = [im.cell(r, 37).value for r in range(2, 6)]
ok("registro: righe 1-3 «Nuova»", esiti[:3] == ["Nuova"] * 3, esiti)
ok("registro: data e ora (colonna Data e ora)", im["AD2"].value == datetime.datetime(2026, 10, 22, 9, 5), im["AD2"].value)
ok("registro: dipendenti numerico", im["AH2"].value == 1200, im["AH2"].value)
ok("registro: URL con accenti e \\/ ripulito", im["AF2"].value == ROWS[0][1], im["AF2"].value)
ok("registro: nome azienda con accenti", im["AE2"].value == ROWS[0][0], im["AE2"].value)
ok("registro: città con spazio («Reggio Calabria»)", im["AJ2"].value == "Milano, Reggio Calabria", im["AJ2"].value)
ok("registro: notte del cambio d'ora (29/03 02:30) conservata", im["AD4"].value == datetime.datetime(2026, 3, 29, 2, 30), im["AD4"].value)
print("Importazione CSV con impostazioni italiane — motore: LibreOffice headless (Excel non disponibile)")
print("\n".join(LOG))
print("RISULTATO: {} superati, {} falliti".format(P, F))
