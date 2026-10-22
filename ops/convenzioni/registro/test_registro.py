"""Prove del registro (revisione 2) con LibreOffice come motore di calcolo.

Uso: python3 test_registro.py <registro.xlsx> <cartella lavoro>
Ogni scenario scrive dati FITTIZI in una copia, la ricalcola con LibreOffice headless e legge i
valori calcolati. Il file consegnato non viene mai modificato.
"""
import sys, os, shutil, subprocess, datetime, json
import openpyxl

SRC, WORK = sys.argv[1], sys.argv[2]
SOFF = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
os.makedirs(WORK, exist_ok=True)
P = F = 0
LOG = []


def ok(desc, cond, got=""):
    global P, F
    if cond:
        P += 1
        LOG.append("OK    | " + desc)
    else:
        F += 1
        LOG.append("FAIL  | {} → {}".format(desc, got))


HEAD = ["Nome azienda", "Sito web azienda (facoltativo)", "Email di riferimento", "Numero dipendenti",
        "Settore dell'azienda", "Città di interesse per la convenzione (una o più)", "Privacy Policy", "ID richiesta",
        "Form Name (ID)", "Submission ID", "Created At", "User ID", "User Agent", "User IP", "Referrer"]


def row(sid, conv="CONV-20261022-AAAA0001", az="Azienda Prova", sito="https:\\/\\/www.azienda.it", email="a@prova.invalid",
        dip="120", sett="Logistica", citta="Milano, Roma", created="2026-10-22 10:15:30"):
    return [az, sito, email, dip, sett, citta, "on", conv, "Convenzioni aziendali (9fc0788)", sid, created, "0",
            "Mozilla/5.0 (prova)", "203.0.113.7", "https://temakinho.com/convenzioni-aziendali/"]


def scenario(name, csv_rows=None, header=None, richieste=None):
    wb = openpyxl.load_workbook(SRC)
    im = wb["Importa CSV"]
    if csv_rows is not None:
        for j, h in enumerate(header or HEAD, start=1):
            im.cell(1, j, h)
        for i, r in enumerate(csv_rows, start=2):
            for j, v in enumerate(r, start=1):
                im.cell(i, j, v)
    if richieste:
        rq = wb["Richieste"]
        tab = rq.tables["tblRichieste"]
        calc = {c.name: c.calculatedColumnFormula.attr_text for c in tab.tableColumns if c.calculatedColumnFormula is not None}
        heads = [c.value for c in rq[1]]
        for i, rr in enumerate(richieste, start=2):
            for k, v in rr.items():
                rq.cell(i, heads.index(k) + 1, v)
            for k, f in calc.items():
                # stessa formula della colonna calcolata, sulla riga i (in Excel l'estensione e' automatica)
                rq.cell(i, heads.index(k) + 1, "=" + f.replace("$I2", "$I{}".format(i)).replace("$A2", "$A{}".format(i)).replace("$S2", "$S{}".format(i)))
        tab.ref = "A1:{}{}".format(openpyxl.utils.get_column_letter(len(heads)), 1 + len(richieste))
    src = os.path.join(WORK, name + ".xlsx")
    wb.save(src)
    out = os.path.join(WORK, "calc")
    os.makedirs(out, exist_ok=True)
    subprocess.run([SOFF, "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", out, src], check=True, capture_output=True)
    return openpyxl.load_workbook(os.path.join(out, name + ".xlsx"), data_only=True)


def errors_in(wb):
    bad = []
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper():
                    bad.append("{}!{}={}".format(ws.title, c.coordinate, c.value))
    return bad


def ak(ws, n):
    return [ws.cell(r, 37).value for r in range(2, 2 + n)]


# S1 — file consegnato, vuoto
w = scenario("s1_vuoto")
im = w["Importa CSV"]
ok("vuoto: nessun errore di formula in tutto il file", not errors_in(w), errors_in(w)[:5])
ok("vuoto: intestazioni in attesa", str(im["AO12"].value).startswith("In attesa"), im["AO12"].value)
ok("vuoto: capacità OK", str(im["AO13"].value).startswith("OK"), im["AO13"].value)
ok("vuoto: nessuna riga calcolata", all(v in (None, "") for v in ak(im, 2000)))
ok("vuoto: area di incolla vuota", im["AO19"].value == "vuota", im["AO19"].value)
rq = w["Richieste"]
ok("vuoto: riga vuota della tabella senza Stato", rq["T2"].value in (None, ""), rq["T2"].value)

# S2 — importazione valida (valori come testo, formato dell'export Elementor)
rows = [row(101), row(102, conv="CONV-20261022-AAAA0002", citta="Reggio Calabria", dip="1"),
        row(103, conv="CONV-20261022-AAAA0003", citta="Lione, Rimini, Bologna", sito=""), row(104, citta="Roma"),
        row(105, created="2026-01-05 08:00:00", dip=" 2500 ")]
w = scenario("s2_valida", rows)
im = w["Importa CSV"]
ok("valida: nessun errore di formula", not errors_in(w), errors_in(w)[:5])
ok("valida: intestazioni OK", im["AO12"].value == "OK", im["AO12"].value)
ok("valida: 5 righe Nuova", ak(im, 5) == ["Nuova"] * 5, ak(im, 5))
ok("valida: conteggio Nuova = 5", im["AO15"].value == 5, im["AO15"].value)
ok("valida: Submission ID numerico", all(isinstance(im.cell(r, 28).value, (int, float)) for r in range(2, 7)), [im.cell(r, 28).value for r in range(2, 7)])
d = im["AD2"].value
ok("valida: data e ora come valore data", isinstance(d, datetime.datetime) and d == datetime.datetime(2026, 10, 22, 10, 15, 30), d)
ok("valida: dipendenti numerico (anche con spazi)", im["AH2"].value == 120 and im["AH6"].value == 2500, (im["AH2"].value, im["AH6"].value))
ok("valida: URL con \\/ ripulito", im["AF2"].value == "https://www.azienda.it", im["AF2"].value)
ok("valida: ID richiesta leggibile", im["AC2"].value == "CONV-20261022-AAAA0001", im["AC2"].value)
ok("valida: nessun IP/browser nelle colonne calcolate", all("203.0.113.7" not in str(im.cell(r, c).value) and "Mozilla" not in str(im.cell(r, c).value) for r in range(2, 7) for c in range(28, 38)))
ok("valida: promemoria per svuotare l'area di incolla", str(im["AO19"].value).startswith("NON vuota"), im["AO19"].value)

# S3 — valori gia' tipizzati da Excel (numeri e date veri)
rows = [row(201, dip=120, created=datetime.datetime(2026, 10, 22, 9, 0, 0)), row(202, dip=3, created=datetime.datetime(2026, 9, 1, 18, 30, 0))]
rows[0][9] = 201
rows[1][9] = 202
w = scenario("s3_tipizzata", rows)
im = w["Importa CSV"]
ok("tipizzata: 2 righe Nuova", ak(im, 2) == ["Nuova"] * 2, ak(im, 2))
ok("tipizzata: data conservata", im["AD2"].value == datetime.datetime(2026, 10, 22, 9, 0, 0), im["AD2"].value)
ok("tipizzata: dipendenti numerico", im["AH3"].value == 3, im["AH3"].value)

# S4 — intestazione Submission ID mancante
h = list(HEAD)
h[9] = "Submission"
w = scenario("s4_senza_id", [row(301), row(302)], header=h)
im = w["Importa CSV"]
ok("senza Submission ID: errore globale", str(im["AO12"].value).startswith("ERRORE") and "Submission ID" in str(im["AO12"].value), im["AO12"].value)
ok("senza Submission ID: importazione bloccata", im["AO14"].value == "BLOCCATA", im["AO14"].value)
ok("senza Submission ID: nessuna riga Nuova", "Nuova" not in ak(im, 2), ak(im, 2))
ok("senza Submission ID: righe marcate bloccate", all(str(v).startswith("Bloccata") for v in ak(im, 2)), ak(im, 2))

# S5 — etichetta di un campo cambiata
h = list(HEAD)
h[3] = "Dipendenti"
w = scenario("s5_etichetta", [row(401)], header=h)
im = w["Importa CSV"]
ok("etichetta cambiata: errore che la nomina", "Numero dipendenti" in str(im["AO12"].value), im["AO12"].value)
ok("etichetta cambiata: bloccata", im["AO14"].value == "BLOCCATA", im["AO14"].value)

# S6 — errori di riga
rows = [row(""), row("abc"), row(503, dip="0"), row(504, dip="1.5"), row(505, dip="abc"), row(506, dip=""),
        row(507, citta=""), row(508, citta="Napoli"), row(509, citta="Milano, Napoli"), row(510)]
w = scenario("s6_righe", rows)
im = w["Importa CSV"]
exp = ["ERRORE: Submission ID mancante", "ERRORE: Submission ID non numerico", "ERRORE: numero dipendenti non valido",
       "ERRORE: numero dipendenti non valido", "ERRORE: numero dipendenti non valido", "ERRORE: numero dipendenti non valido",
       "ERRORE: nessuna città", "ERRORE: città non prevista", "ERRORE: città non prevista", "Nuova"]
got = ak(im, 10)
for i, (e, g) in enumerate(zip(exp, got)):
    ok("riga {}: {}".format(i + 2, e), e == g, g)
ok("errori di riga: conteggio errori = 9", im["AO18"].value == 9, im["AO18"].value)
ok("errori di riga: importazione non bloccata", im["AO14"].value == "OK", im["AO14"].value)

# S7 — doppioni e codice CONV non affidabile
rows = [row(601, conv="CONV-X"), row(601, conv="CONV-X"), row(602, conv="CONV-X"), row(603, conv="")]
w = scenario("s7_doppioni", rows)
im = w["Importa CSV"]
ok("stesso Submission ID due volte: secondo = doppione", ak(im, 2) == ["Nuova", "Doppione nell'export"], ak(im, 2))
ok("stesso CONV con Submission ID diverso: tenuto", ak(im, 3)[2] == "Nuova", ak(im, 3))
ok("riga senza CONV ma con Submission ID: tenuta", ak(im, 4)[3] == "Nuova", ak(im, 4))

# S8 — reimportazione con registro popolato: stato e note preservati
reg = [
    {"ID archivio WordPress": 101, "ID richiesta": "CONV-A", "Nome azienda": "Azienda Uno", "Città selezionate": "Roma, Rimini", "Stato (manuale)": "In valutazione", "Note interne": "Richiamare lunedì"},
    {"ID archivio WordPress": 102, "ID richiesta": "CONV-B", "Nome azienda": "Azienda Due", "Città selezionate": "Reggio Calabria", "Note interne": "Nota fittizia"},
]
w = scenario("s8_reimport", [row(101), row(102), row(103)], richieste=reg)
im, rq = w["Importa CSV"], w["Richieste"]
ok("reimport: righe gia' presenti riconosciute", ak(im, 3) == ["Già nel registro", "Già nel registro", "Nuova"], ak(im, 3))
ok("reimport: solo la nuova da copiare", im["AO15"].value == 1, im["AO15"].value)
ok("registro: stato manuale conservato", rq["T2"].value == "In valutazione", rq["T2"].value)
ok("registro: riga senza stato = Da valutare", rq["T3"].value == "Da valutare", rq["T3"].value)
ok("registro: note interne conservate", rq["U2"].value == "Richiamare lunedì" and rq["U3"].value == "Nota fittizia", (rq["U2"].value, rq["U3"].value))
heads = [c.value for c in rq[1]]
city = lambda r, c: rq.cell(r, heads.index(c) + 1).value
ok("registro: filtro città Roma sulla riga «Roma, Rimini»", city(2, "Roma") == "Sì" and city(2, "Rimini") == "Sì" and city(2, "Milano") in (None, ""), [city(2, c) for c in ("Milano", "Roma", "Rimini")])
ok("registro: Reggio Calabria riconosciuta", city(3, "Reggio Calabria") == "Sì" and city(3, "Roma") in (None, ""), [city(3, c) for c in ("Roma", "Reggio Calabria")])

# S9 — capacita'
rows = [row(10000 + i) for i in range(2000)]
w = scenario("s9_2000", rows)
im = w["Importa CSV"]
ok("capacità: 2000 righe accettate", str(im["AO13"].value).startswith("OK") and im["AO15"].value == 2000, (im["AO13"].value, im["AO15"].value))
rows.append(row(99999))
w = scenario("s9_2001", rows)
im = w["Importa CSV"]
ok("capacità: 2001 righe → errore esplicito", str(im["AO13"].value).startswith("ERRORE"), im["AO13"].value)
ok("capacità superata: importazione bloccata", im["AO14"].value == "BLOCCATA" and im["AO15"].value == 0, (im["AO14"].value, im["AO15"].value))

# S10 — struttura del file consegnato
wb = openpyxl.load_workbook(SRC)
rq, im = wb["Richieste"], wb["Importa CSV"]
t = rq.tables.get("tblRichieste")
ok("tabella tblRichieste presente", t is not None)
ok("chiave = prima colonna «ID archivio WordPress»", rq["A1"].value == "ID archivio WordPress", rq["A1"].value)
calcnames = sorted(c.name for c in t.tableColumns if c.calculatedColumnFormula is not None)
ok("colonne calcolate della tabella (6 città + Stato)", calcnames == sorted(["Milano", "Roma", "Bologna", "Lione", "Rimini", "Reggio Calabria", "Stato"]), calcnames)
ok("elenco stati sulla colonna Stato (manuale)", any("Elenchi!$A$2" in (dv.formula1 or "") for dv in rq.data_validations.dataValidation))
ok("evidenziazione degli ID archivio ripetuti", len(rq.conditional_formatting) >= 1)
ok("formule di importazione fino alla riga 2001", im["AK2001"].value is not None and im["AK2002"].value is None)
ok("nessun dato nell'area di incolla", all(im.cell(r, c).value is None for r in range(1, 2002) for c in range(1, 27)))
ok("registro senza righe di dati", all(rq.cell(2, c).value is None or str(rq.cell(2, c).value).startswith("=") for c in range(1, 22)))
ok("nessuna colonna IP o browser nel registro", not any(h and ("IP" in str(h) or "Agent" in str(h)) for h in [c.value for c in rq[1]]))

print("\n".join(LOG))
print("RISULTATO: {} superati, {} falliti".format(P, F))
