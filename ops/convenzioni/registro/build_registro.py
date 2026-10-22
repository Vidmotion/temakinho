"""Genera il registro operativo delle richieste di convenzione (revisione 2).

Chiave autorevole: Submission ID dell'archivio WordPress (Elementor Submissions). Il codice
CONV-... resta come riferimento leggibile. Importazione manuale da CSV, capacita' 2000 righe,
errori espliciti, nessun dato nel file consegnato.
Uso: python3 build_registro.py <file.xlsx>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter

OUT = sys.argv[1]
CAP = 2000                      # righe di dati per importazione
LAST = CAP + 1                  # ultima riga con formule nel foglio Importa CSV
CITIES = ["Milano", "Roma", "Bologna", "Lione", "Rimini", "Reggio Calabria"]
STATES = ["Da valutare", "In valutazione", "Convenzionata", "Non accolta"]
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="F7F1E3")

wb = Workbook()

# ------------------------------------------------------------------ Elenchi
el = wb.active
el.title = "Elenchi"
el["A1"], el["B1"] = "Stati", "Città"
for i, s in enumerate(STATES, start=2):
    el.cell(i, 1, s)
for i, c in enumerate(CITIES, start=2):
    el.cell(i, 2, c)
el["A1"].font = el["B1"].font = BOLD
el.column_dimensions["A"].width = 18
el.column_dimensions["B"].width = 18

# ------------------------------------------------------------------ Richieste
rq = wb.create_sheet("Richieste", 0)
cols = (["ID archivio WordPress", "ID richiesta", "Data e ora", "Nome azienda", "Sito web",
         "Email di riferimento", "Numero dipendenti", "Settore", "Città selezionate"]
        + CITIES
        + ["Newsletter: adesione", "Newsletter: data consenso", "Newsletter: versione testo",
           "Stato (manuale)", "Stato", "Note interne"])
for i, name in enumerate(cols, start=1):
    c = rq.cell(1, i, name)
    c.font = BOLD
    c.fill = HEAD
    c.alignment = Alignment(wrap_text=True, vertical="top")
ncol = len(cols)
colof = {name: get_column_letter(i) for i, name in enumerate(cols, start=1)}
I = colof["Città selezionate"]
S = colof["Stato (manuale)"]
A = colof["ID archivio WordPress"]
calc = {}
for city in CITIES:
    L = colof[city]
    # confronto per voce intera, indipendente da spazi: ",Milano," dentro ",Milano,Roma,"
    calc[city] = ('IF(${I}2="","",IF(ISNUMBER(SEARCH(","&SUBSTITUTE({L}$1," ","")&",",'
                  '","&SUBSTITUTE(${I}2," ","")&",")),"Sì",""))').format(I=I, L=L)
calc["Stato"] = 'IF(${A}2="","",IF(${S}2="","Da valutare",${S}2))'.format(A=A, S=S)
for name, f in calc.items():
    rq["{}2".format(colof[name])] = "=" + f
ref = "A1:{}2".format(get_column_letter(ncol))
tab = Table(displayName="tblRichieste", ref=ref)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
tab.autoFilter = AutoFilter(ref=ref)
# Colonne costruite a mano: i nomi DEVONO coincidere con le intestazioni (altrimenti Excel "ripara"
# il file) e le colonne calcolate portano la formula, che Excel estende alle nuove righe.
for i, name in enumerate(cols, start=1):
    tc = TableColumn(id=i, name=name)
    if name in calc:
        tc.calculatedColumnFormula = TableFormula(attr_text=calc[name])
    tab.tableColumns.append(tc)
rq.add_table(tab)
rq.freeze_panes = "B2"
widths = {"ID archivio WordPress": 12, "ID richiesta": 24, "Data e ora": 17, "Nome azienda": 28,
          "Sito web": 28, "Email di riferimento": 28, "Numero dipendenti": 11, "Settore": 24,
          "Città selezionate": 28, "Stato (manuale)": 16, "Stato": 14, "Note interne": 40}
for name, L in colof.items():
    rq.column_dimensions[L].width = widths.get(name, 11)
MAXR = 20001
rq.column_dimensions[colof["Data e ora"]].number_format = "dd/mm/yyyy hh:mm"
for name, fmt in (("Data e ora", "dd/mm/yyyy hh:mm"), ("Numero dipendenti", "0"),
                  ("ID archivio WordPress", "0"), ("Newsletter: data consenso", "dd/mm/yyyy hh:mm")):
    for r in range(2, 3):
        rq["{}{}".format(colof[name], r)].number_format = fmt
dv = DataValidation(type="list", formula1="Elenchi!$A$2:$A${}".format(len(STATES) + 1), allow_blank=True,
                    showErrorMessage=True, errorTitle="Stato non previsto", error="Scegli uno stato dall'elenco.")
rq.add_data_validation(dv)
dv.add("{}2:{}{}".format(S, S, MAXR))
red = PatternFill("solid", fgColor="F4C7C3")
rq.conditional_formatting.add("{A}2:{A}{m}".format(A=A, m=MAXR),
                              FormulaRule(formula=['AND(${A}2<>"",COUNTIF(${A}:${A},${A}2)>1)'.format(A=A)], fill=red))

# ------------------------------------------------------------------ Importa CSV
im = wb.create_sheet("Importa CSV", 1)
# Mappa intestazioni del CSV Elementor -> campo; AO = colonna trovata (0 = assente)
MAP = [  # riga, intestazione CSV, campo, obbligatoria
    (2, "Submission ID", "ID archivio WordPress (chiave)", "Sì"),
    (3, "Created At", "Data e ora", "Sì"),
    (4, "ID richiesta", "ID richiesta (CONV)", "No"),
    (5, "Nome azienda", "Nome azienda", "Sì"),
    (6, "Sito web azienda (facoltativo)", "Sito web", "Sì"),
    (7, "Email di riferimento", "Email di riferimento", "Sì"),
    (8, "Numero dipendenti", "Numero dipendenti", "Sì"),
    (9, "Settore dell'azienda", "Settore", "Sì"),
    (10, "Città di interesse per la convenzione (una o più)", "Città selezionate", "Sì"),
]
for h, v in (("AN1", "Intestazione nel CSV Elementor"), ("AO1", "Colonna trovata"), ("AP1", "Campo nel registro"), ("AQ1", "Obbligatoria")):
    im[h] = v
    im[h].font = BOLD
    im[h].fill = HEAD
for r, head, field, req in MAP:
    im["AN{}".format(r)] = head
    im["AO{}".format(r)] = "=IFERROR(MATCH(AN{r},$A$1:$Z$1,0),0)".format(r=r)
    im["AP{}".format(r)] = field
    im["AQ{}".format(r)] = req
missing = "&".join('IF(AND($AQ${r}="Sì",$AO${r}=0),$AN${r}&"; ","")'.format(r=r) for r, *_ in MAP)
ctrl = [
    (12, "Intestazioni", '=IF(COUNTA($A$1:$Z$1)=0,"In attesa: incolla il CSV in A1",IF(COUNTIFS($AO$2:$AO$10,0,$AQ$2:$AQ$10,"Sì")>0,"ERRORE: intestazioni mancanti o cambiate: "&' + missing + ',"OK"))'),
    (13, "Capacità", '=IF(COUNTA($A${n}:$Z$200000)>0,"ERRORE: più di {cap} righe di dati: importa l\'export in più parti","OK (fino a {cap} righe)")'.format(n=LAST + 1, cap=CAP)),
    (14, "Importazione", '=IF(OR(LEFT($AO$12,6)="ERRORE",LEFT($AO$13,6)="ERRORE"),"BLOCCATA","OK")'),
    (15, "Righe «Nuova» da copiare", '=COUNTIF($AK$2:$AK${l},"Nuova")'.format(l=LAST)),
    (16, "Già nel registro", '=COUNTIF($AK$2:$AK${l},"Già nel registro")'.format(l=LAST)),
    (17, "Doppioni nell'export", '=COUNTIF($AK$2:$AK${l},"Doppione nell\'export")'.format(l=LAST)),
    (18, "Righe con errori", '=COUNTIF($AK$2:$AK${l},"ERRORE*")+COUNTIF($AK$2:$AK${l},"Bloccata*")'.format(l=LAST)),
    (19, "Area di incolla (A:Z)", '=IF(COUNTA($A$1:$Z${n})>0,"NON vuota: contiene IP e browser, svuotala a importazione finita","vuota")'.format(n=LAST + 50)),
]
for r, label, f in ctrl:
    im["AN{}".format(r)] = label
    im["AN{}".format(r)].font = BOLD
    im["AO{}".format(r)] = f
im["AM1"] = "CONTROLLI →"
im["AM1"].font = BOLD
heads = ["ID archivio WordPress", "ID richiesta", "Data e ora", "Nome azienda", "Sito web",
         "Email di riferimento", "Numero dipendenti", "Settore", "Città selezionate", "Esito controllo"]
for i, h in enumerate(heads):
    c = im.cell(1, 28 + i, h)  # AB=28
    c.font = BOLD
    c.fill = HEAD


def raw(row, k):
    return 'INDEX($A{r}:$Z{r},$AO${k})'.format(r=row, k=k)


residual = ',"&SUBSTITUTE($AJ{r}," ","")&","'
for row in range(2, LAST + 1):
    empty = 'COUNTA($A{r}:$Z{r})=0'.format(r=row)
    # INDEX su una cella vuota restituisce 0, non "": i campi di testo si forzano con &""
    g = lambda k: 'IF($AO${k}>0,{x}&"","")'.format(k=k, x=raw(row, k))
    im["AB{}".format(row)] = '=IF({e},"",IF($AO$2=0,"",IF({x}&""="","",IFERROR(VALUE({x}&""),{x}&""))))'.format(e=empty, x=raw(row, 2))
    im["AC{}".format(row)] = '=IF({e},"",{v})'.format(e=empty, v=g(4))
    d = raw(row, 3)
    im["AD{}".format(row)] = ('=IF({e},"",IF($AO$3=0,"",IF(ISNUMBER({d}),{d},IFERROR(DATE(VALUE(LEFT({d}&"",4)),VALUE(MID({d}&"",6,2)),VALUE(MID({d}&"",9,2)))'
                              '+TIME(VALUE(MID({d}&"",12,2)),VALUE(MID({d}&"",15,2)),VALUE(MID({d}&"",18,2))),{d}&""))))').format(e=empty, d=d)
    im["AE{}".format(row)] = '=IF({e},"",{v})'.format(e=empty, v=g(5))
    im["AF{}".format(row)] = '=IF({e},"",SUBSTITUTE({v},"\\/","/"))'.format(e=empty, v=g(6))
    im["AG{}".format(row)] = '=IF({e},"",{v})'.format(e=empty, v=g(7))
    im["AH{}".format(row)] = '=IF({e},"",IF($AO$8=0,"",IF(ISNUMBER({x}),{x},IFERROR(VALUE(TRIM({x}&"")),TRIM({x}&"")))))'.format(e=empty, x=raw(row, 8))
    im["AI{}".format(row)] = '=IF({e},"",{v})'.format(e=empty, v=g(9))
    im["AJ{}".format(row)] = '=IF({e},"",{v})'.format(e=empty, v=g(10))
    # residuo dopo aver tolto ogni citta' dell'elenco: se resta qualcosa, c'e' una citta' non prevista
    s = '","&SUBSTITUTE($AJ{r}," ","")&","'.format(r=row)
    for i in range(2, 2 + len(CITIES)):
        s = 'SUBSTITUTE({s},","&SUBSTITUTE(Elenchi!$B${i}," ","")&",",",")'.format(s=s, i=i)
    unknown = 'LEN(SUBSTITUTE({s},",",""))>0'.format(s=s)
    im["AK{}".format(row)] = (
        '=IF({e},"",IF($AO$14="BLOCCATA","Bloccata: vedi i controlli in alto",'
        'IF($AB{r}="","ERRORE: Submission ID mancante",'
        'IF(NOT(ISNUMBER($AB{r})),"ERRORE: Submission ID non numerico",'
        'IF(NOT(ISNUMBER($AH{r})),"ERRORE: numero dipendenti non valido",'
        'IF(OR($AH{r}<1,$AH{r}<>INT($AH{r})),"ERRORE: numero dipendenti non valido",'
        'IF(TRIM($AJ{r})="","ERRORE: nessuna città",'
        'IF({u},"ERRORE: città non prevista",'
        'IF(COUNTIF(Richieste!$A:$A,$AB{r})>0,"Già nel registro",'
        'IF(COUNTIF($AB$2:$AB{r},$AB{r})>1,"Doppione nell\'export","Nuova"))))))))))'
    ).format(e=empty, r=row, u=unknown)
    im["AB{}".format(row)].number_format = "0"
    im["AD{}".format(row)].number_format = "dd/mm/yyyy hh:mm"
    im["AH{}".format(row)].number_format = "0"
for L, w in (("AB", 12), ("AC", 24), ("AD", 17), ("AE", 26), ("AF", 26), ("AG", 26), ("AH", 11), ("AI", 20), ("AJ", 26), ("AK", 34), ("AN", 44), ("AO", 34), ("AP", 28)):
    im.column_dimensions[L].width = w
im.freeze_panes = "A2"

# ------------------------------------------------------------------ Istruzioni
ins = wb.create_sheet("Istruzioni", 0)
lines = [
    ("Registro richieste di convenzione – Temakinho (copia operativa, revisione 2)", True),
    ("", False),
    ("REGISTRO MANUALE – nessuna sincronizzazione automatica attiva.", True),
    ("La fonte persistente è l'archivio WordPress: Elementor › Submissions › modulo \"Convenzioni aziendali\" (solo amministratori). Questo file si aggiorna solo con l'importazione descritta sotto.", False),
    ("", False),
    ("Struttura", True),
    ("Foglio Richieste: tabella \"tblRichieste\", una riga per invio. Chiave = ID archivio WordPress (il Submission ID dell'export): è l'unico identificativo affidabile. ID richiesta (CONV-…) serve solo come riferimento leggibile e può ripetersi. Gli ID archivio ripetuti si colorano di rosso.", False),
    ("Le colonne città e Stato sono calcolate e si estendono da sole alle nuove righe della tabella.", False),
    ("Stato (manuale) e Note interne le compila solo il team: l'importazione aggiunge righe e non modifica mai quelle esistenti.", False),
    ("", False),
    ("Importazione manuale dall'archivio WordPress", True),
    ("1. Elementor › Submissions › filtra \"Convenzioni aziendali\" › Export All to CSV.", False),
    ("2. In Excel: Dati › Da testo/CSV, origine 65001 Unicode (UTF-8), delimitatore virgola. Copia tutto, intestazioni comprese.", False),
    ("3. Foglio Importa CSV: incolla dalla cella A1 (area A:Z). Le colonne AB:AK si calcolano da sole; i controlli sono in AN:AO.", False),
    ("4. Se Importazione = BLOCCATA, correggi la causa indicata (intestazioni o capacità) e reincolla. Le righe con ERRORE non vanno copiate: correggi il dato in WordPress o annotalo.", False),
    ("5. Filtra \"Esito controllo\" = Nuova, copia AB:AJ delle righe visibili e incollale come VALORI nella prima riga libera di tblRichieste (colonne da ID archivio WordPress a Città selezionate).", False),
    ("6. Svuota l'area A:Z del foglio Importa CSV: il CSV di Elementor contiene IP e browser, che non vanno conservati qui. Il controllo \"Area di incolla\" lo ricorda.", False),
    ("Le righe \"Già nel registro\" e \"Doppione nell'export\" non vanno copiate: così stato e note non vengono mai sovrascritti e ogni invio compare una sola volta.", False),
    ("", False),
    ("Capacità", True),
    ("Fino a {} righe di dati per importazione. Oltre, il controllo Capacità blocca l'importazione: esporta per periodo (filtro data in Submissions) e importa in più parti, oppure estendi le formule AB:AK copiando l'ultima riga verso il basso e aggiorna il controllo Capacità (cella AO13) con la nuova prima riga libera.".format(CAP), False),
    ("Se le etichette dei campi del modulo cambiano, aggiorna la colonna AN del foglio Importa CSV. Se cambia l'elenco delle città, aggiorna il foglio Elenchi, aggiungi la colonna città in tblRichieste e la voce nella formula di controllo (colonna AK).", False),
    ("", False),
    ("Stati", True),
    ("Ogni nuova richiesta è \"Da valutare\": il modulo non attiva convenzioni né riconosce condizioni. \"Convenzionata\" si imposta solo a mano quando la convenzione è davvero approvata.", False),
    ("", False),
    ("Futura sincronizzazione (non attiva)", True),
    ("Un collegamento automatico dovrà solo aggiungere righe a tblRichieste usando l'ID archivio WordPress come chiave (controllo di esistenza prima dell'inserimento), senza toccare Stato (manuale) e Note interne. Gli invii non sincronizzati si recuperano con il foglio Importa CSV.", False),
    ("", False),
    ("Newsletter", True),
    ("Il modulo attuale non raccoglie l'adesione alla newsletter: le colonne restano predisposte per un'attivazione futura, facoltativa e senza selezione preventiva.", False),
    ("", False),
    ("Condivisione", True),
    ("Contiene dati di contatto: salvalo in un'area aziendale Temakinho con accesso nominativo, senza link pubblici.", False),
]
for i, (t, b) in enumerate(lines, start=1):
    c = ins.cell(i, 1, t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if b:
        c.font = BOLD
ins.column_dimensions["A"].width = 140
wb.active = 0
wb.properties.creator = "Temakinho"
wb.properties.title = "Registro richieste di convenzione"
wb.save(OUT)
print("salvato", OUT)
